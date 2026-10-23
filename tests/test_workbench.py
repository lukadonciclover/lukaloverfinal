from io import BytesIO
import math

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from equitylens.models import DCFInputs
from equitylens.workbench import (
    Assets,
    DDM,
    ForecastAssumptions,
    Peer,
    WACC,
    Workbench,
    calculate_assets,
    calculate_ddm,
    calculate_wacc,
    evaluate,
    normalized_history,
)
from equitylens.workbook import export_workbench

client = TestClient(app)


def model(horizon=5):
    return DCFInputs(
        base_revenue=100,
        revenue_growth=[0] * horizon,
        ebit_margin=[0.1] * horizon,
        tax_rate=[0] * horizon,
        depreciation_pct_revenue=[0] * horizon,
        capex_pct_revenue=[0] * horizon,
        nwc_pct_revenue=[0] * horizon,
        wacc=0.1,
        terminal_growth=0,
    )


def project(horizon=5):
    balances = dict(
        revenue=100,
        cash=20,
        total_assets=200,
        total_liabilities=80,
        total_equity=120,
        short_term_debt=10,
        long_term_debt=40,
        net_working_capital=15,
        property_plant_equipment=100,
        ebitda=20,
        ebit=17,
        net_income=10,
    )
    figures = [
        dict(
            company_id=1,
            metric=k,
            period="2025",
            period_start="2025-01-01", period_end="2025-12-31", fiscal_calendar="test-calendar",
            value=v,
            currency="USD",
            unit_scale=1000,
        )
        for k, v in balances.items()
    ]
    return Workbench(
        currency="USD",
        figures=figures,
        forecast=ForecastAssumptions(horizon=horizon),
        dcf=model(horizon),
        cash=20000,
        debt=50000,
        shares=1000,
        price=50,
        wacc=WACC(
            risk_free=0.04,
            beta=1.2,
            equity_risk_premium=0.05,
            cost_of_debt=0.06,
            tax_rate=0.25,
            market_equity=80000,
            market_debt=20000,
        ),
        ddm=DDM(
            dividend=2,
            growth=0.05,
            terminal_growth=0.02,
            cost_of_equity=0.1,
            horizon=horizon,
        ),
        assets=Assets(
            book_assets=200000, book_liabilities=80000, asset_adjustment=10000
        ),
        peers=[
            Peer(
                company="Peer",
                reporting_basis="annual", basis_reviewed=True, financial_period="2025", period_end="2025-12-31", measurement_date="2026-01-01",
                currency="USD",
                unit_scale=1000,
                market_cap=150,
                debt=50,
                cash=20,
                revenue=100,
                ebitda=20,
                ebit=17,
                net_income=10,
                book_equity=120,
            )
        ],
        transactions=[
            Peer(
                company="Target",
                reporting_basis="annual", basis_reviewed=True, financial_period="2025", period_end="2025-12-31", measurement_date="2026-01-01",
                currency="USD",
                unit_scale=1000,
                acquisition_date="2026-01-01",
                enterprise_value=200,
                revenue=100,
                ebitda=20,
                ebit=17,
            )
        ],
    )


@pytest.mark.parametrize("horizon", [3, 5, 10])
def test_integrated_forecast_balances_and_fcff_matches_dcf(horizon):
    output = evaluate(project(horizon))
    assert output["errors"] == {}
    assert len(output["forecast"]) == horizon
    for forecast, dcf in zip(output["forecast"], output["models"]["DCF"]["forecast"]):
        assert forecast["balance_check"] == pytest.approx(0, abs=1e-7)
        assert forecast["fcff"] == pytest.approx(dcf["fcff"])


def test_wacc_ddm_assets_known_answers():
    w = WACC(
        risk_free=0.04,
        beta=1.2,
        equity_risk_premium=0.05,
        cost_of_debt=0.06,
        tax_rate=0.25,
        market_equity=80,
        market_debt=20,
    )
    assert calculate_wacc(w)["wacc"] == pytest.approx(0.089)
    assert calculate_ddm(
        DDM(
            dividend=2, growth=0.02, terminal_growth=0.02, cost_of_equity=0.1, horizon=0
        )
    )["per_share"] == pytest.approx(25.5)
    assert calculate_assets(Assets(book_assets=10, book_liabilities=20), None) == {
        "book_equity": -10,
        "equity_value": -10,
        "per_share": None,
    }
    with pytest.raises(ValueError):
        calculate_ddm(
            DDM(
                dividend=2,
                growth=0.1,
                terminal_growth=0.1,
                cost_of_equity=0.1,
                horizon=0,
            )
        )
    with pytest.raises(ValueError):
        calculate_wacc(w.model_copy(update={"market_equity": 0, "market_debt": 0}))


def test_dcf_timing_negative_cash_flow_and_missing_shares():
    p = Workbench(currency="USD", dcf=model(), cash=20, debt=30, shares=10)
    r = evaluate(p)["models"]["DCF"]
    assert r["enterprise_value"] == pytest.approx(100)
    assert r["equity_value"] == pytest.approx(90)
    assert r["per_share"] == pytest.approx(9)
    p.dcf.ebit_margin = [-0.1] * 5
    p.dcf.tax_rate = [0.25] * 5
    p.shares = None
    r = evaluate(p)["models"]["DCF"]
    assert r["forecast"][0]["cash_taxes"] == 0
    assert r["equity_value"] < 0 and r["per_share"] is None


def test_models_are_independent_and_missing_data_is_explicit():
    p = project()
    p.figures = []
    r = evaluate(p)
    assert "forecast" in r["errors"] and "dcf" in r["errors"]
    assert "DCF" not in r["models"] and "DDM" in r["models"]
    p = project()
    p.cash = None
    assert "dcf" in evaluate(p)["errors"]
    p = project()
    p.dcf.wacc = 0
    p.wacc = None
    assert "DCF" not in evaluate(p)["models"]


def test_scale_currency_peer_denominators_and_selection():
    p = project()
    original = evaluate(p)
    for f in p.figures:
        f.value *= 1000
        f.unit_scale = 1
    assert (
        evaluate(p)["models"]["DCF"]["equity_value"]
        == original["models"]["DCF"]["equity_value"]
    )
    p.currency = "EUR"
    with pytest.raises(ValueError, match="currency"):
        normalized_history(p)
    p = project()
    p.peers[0].ebitda = -1
    assert (
        evaluate(p)["models"]["Trading comparables"]["multiples"][0]["ev_ebitda"]
        is None
    )
    p.peers[0].selected = False
    assert evaluate(p)["models"]["Trading comparables"]["ranges"] == []
    p.peers[0].selected = True
    p.peers[0].currency = "EUR"
    assert "Trading comparables" in evaluate(p)["errors"]


def test_scenarios_sensitivity_and_recalculation():
    p = project()
    r = evaluate(p)
    base = r["models"]["DCF"]["equity_value"]
    assert r["scenarios"]["Base"]["equity_value"] == base
    assert r["sensitivities"]["growth"]["values"][2][2] == base
    assert r["sensitivities"]["operating"]["values"][2][2] == base
    p.forecast.growth = 0.1
    assert evaluate(p)["models"]["DCF"]["equity_value"] != base
    p.dcf.terminal_method = "exit_multiple"
    r = evaluate(p)
    assert (
        r["sensitivities"]["exit_multiple"]["values"][2][2]
        == r["models"]["DCF"]["equity_value"]
    )


def test_api_validation_and_obsolete_routes():
    for path in [
        "prompts/agent",
        "chat/prompt",
        "thesis",
        "evidence/verify",
        "inference/generate",
        "reports/pdf",
    ]:
        assert client.post("/v1/" + path, json={}).status_code == 404
    p = project().model_dump(mode="json")
    assert client.post("/v1/workbench/calculate", json=p).status_code == 200
    p["shares"] = 0
    assert client.post("/v1/workbench/calculate", json=p).status_code == 422
    response = client.post(
        "/v1/workbench/import-peers",
        files={
            "file": (
                "peers.csv",
                b"company,currency,unit_scale,revenue\nPeer,USD,1000,100\n",
                "text/csv",
            )
        },
    )
    assert (
        response.status_code == 200 and response.json()["rows"][0]["selected"] is False
    )


def test_export_formulas_sources_and_injection():
    p = project()
    p.figures[0].notes = '=HYPERLINK("https://example.com")'
    wb = load_workbook(BytesIO(export_workbench(p)))
    assert wb["Historical statements"]["I2"].data_type == "s"
    assert wb["Projected statements"]["B2"].data_type == "f"
    assert wb["DCF"]["T8"].data_type == "f"
    assert wb["DCF"]["B2"].font.color.rgb == "00008000"
    assert wb["DCF"]["I2"].font.color.rgb == "00000000"
    assert wb["Forecast assumptions"]["B3"].font.color.rgb == "000000FF"
    assert wb.calculation.fullCalcOnLoad


def test_excel_recalculation_matches_application(tmp_path):
    import formulas

    p = project(3)
    expected = evaluate(p)
    path = tmp_path / "model.xlsx"
    path.write_bytes(export_workbench(p))
    calculated = formulas.ExcelModel().loads(str(path)).finish().calculate()

    def value(sheet, cell):
        key = next(
            k
            for k in calculated
            if k.upper().replace("'", "").endswith(f"]{sheet.upper()}!{cell.upper()}")
        )
        return float(calculated[key].value[0, 0])

    assert value("DCF", "T8") == pytest.approx(
        expected["models"]["DCF"]["equity_value"]
    )
    assert value("DDM", "F1") == pytest.approx(expected["models"]["DDM"]["per_share"])
    assert value("Asset valuation", "B8") == pytest.approx(
        expected["models"]["Asset based"]["equity_value"]
    )
    assert value("WACC", "B10") == pytest.approx(expected["wacc"]["wacc"])
    for index, row in enumerate(expected["forecast"], 2):
        assert value("Projected statements", f"Z{index}") == pytest.approx(0, abs=1e-6)
    for index, scenario in enumerate(expected["scenarios"].values(), 2):
        assert value("Scenarios", f"G{index}") == pytest.approx(
            scenario["equity_value"]
        )
    index = 2
    for grid in expected["sensitivities"].values():
        for row in grid["values"]:
            for number in row:
                if number is not None:
                    assert value("Sensitivity", f"D{index}") == pytest.approx(number)
                index += 1
    for sheet, key in [
        ("Comparables", "Trading comparables"),
        ("Transactions", "Precedent transactions"),
    ]:
        wb = load_workbook(path)
        ranges = expected["models"][key]["ranges"]
        start = next(
            cell.row + 1
            for cell in wb[sheet]["A"]
            if cell.value == "Implied valuations"
        )
        for index, row in enumerate(ranges, start):
            assert value(sheet, f"D{index}") == pytest.approx(row["equity_value"])
    # Changing an exported driver must recalculate the forecast and valuation.
    wb = load_workbook(path)
    wb["Forecast assumptions"]["B3"] = 0.08
    wb.save(path)
    p.forecast.growth = 0.08
    calculated = formulas.ExcelModel().loads(str(path)).finish().calculate()
    assert value("DCF", "T8") == pytest.approx(
        evaluate(p)["models"]["DCF"]["equity_value"]
    )


@pytest.mark.parametrize("scale", [1, 1000, 1000000])
def test_actual_unit_normalization(scale):
    p = project()
    for f in p.figures:
        f.value = f.value * f.unit_scale / scale
        f.unit_scale = scale
    assert normalized_history(p)["2025"]["revenue"] == 100000


def test_explicit_revolver_financing_and_unfunded_cash_shortfall():
    p = project()
    p.forecast.dividends = 50000
    p.forecast.minimum_cash = 10000
    r = evaluate(p)
    assert r["forecast"][0]["revolver_draw"] > 0
    assert r["forecast"][0]["cash"] == 10000
    assert r["forecast"][0]["balance_check"] == pytest.approx(0)
    p.forecast.financing = "none"
    assert "forecast" in evaluate(p)["errors"]


def test_zero_denominators_bad_dates_and_invalid_balance_sheet():
    p = project()
    p.peers[0].net_income = 0
    assert evaluate(p)["models"]["Trading comparables"]["multiples"][0]["pe"] is None
    p.transactions[0].acquisition_date = "not a date"
    assert "Precedent transactions" in evaluate(p)["errors"]
    next(f for f in p.figures if f.metric == "total_assets").value += 1
    assert "forecast" in evaluate(p)["errors"]
