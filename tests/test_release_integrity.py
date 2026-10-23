from datetime import date
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app
from equitylens.models import FinancialFigure
from equitylens.periods import classify, ttm_figures
from equitylens.workbench import Workbench, Peer, ForwardTarget, evaluate
from equitylens.workbook import export_workbench
from test_workbench import project

client=TestClient(app)
def quarters(metric='revenue'):
    return [FinancialFigure(company_id=1,metric=metric,period=f'2025-Q{i}',value=100,unit_scale=1000,period_start=s,period_end=e,fiscal_calendar='calendar') for i,(s,e) in enumerate([('2025-01-01','2025-03-31'),('2025-04-01','2025-06-30'),('2025-07-01','2025-09-30'),('2025-10-01','2025-12-31')],1)]

def peer(**changes):
    return Peer(company='Acme',currency='USD',unit_scale=1,enterprise_value=1000000,revenue=100000,reporting_basis='annual',financial_period='2025',period_end='2025-12-31',measurement_date='2026-01-01',basis_reviewed=True,**changes)

def test_four_quarter_ttm_has_sources_and_normalizes_without_summing_stocks():
    rows=quarters()+quarters('cash')
    rows[0].value=100000;rows[0].unit_scale=1
    ttm=ttm_figures(rows,'USD')['figures']
    assert len(ttm)==1 and ttm[0]['value']==400000 and len(ttm[0]['inputs'])==4
    assert ttm[0]['period_start']=='2025-01-01' and ttm[0]['period_end']=='2025-12-31'
    assert classify('FY25')=='unsupported'

@pytest.mark.parametrize('fault',['missing','duplicate','overlap','currency','calendar','undated','ytd'])
def test_invalid_quarters_never_produce_ttm(fault):
    rows=quarters()
    if fault=='missing': rows.pop(1)
    if fault=='duplicate': rows.append(rows[1])
    if fault=='overlap': rows[1].period_start='2025-03-15'
    if fault=='currency': rows[1].currency='EUR'
    if fault=='calendar': rows[1].fiscal_calendar='changed'
    if fault=='undated': rows[1].period_start=''
    if fault=='ytd': rows[1].period_start='2025-01-01'
    assert ttm_figures(rows,'USD')['figures']==[]

def test_mixed_frequency_analysis_and_annual_forecast_stay_separate():
    p=project();baseline=evaluate(p)
    p.figures+=quarters()
    output=evaluate(p)
    assert output['forecast']==baseline['forecast']
    assert output['models']['DCF']==baseline['models']['DCF']
    assert output['models']['Trading comparables']['ranges']==baseline['models']['Trading comparables']['ranges']
    data=client.post('/v1/financials/analyze',json={'figures':[f.model_dump() for f in p.figures],'currency':'USD'})
    assert data.status_code==200
    rows={r['period']:r for r in data.json()['analysis']}
    assert rows['2025-Q1']['revenue_growth'] is None
    assert rows['2025-Q2']['revenue_growth']==0

def test_annual_multiples_block_quarter_only_and_ttm_values_are_complete():
    p=Workbench(currency='USD',figures=quarters(),peers=[peer()],cash=0,debt=0)
    assert 'quarterly' in evaluate(p)['errors']['Trading comparables']
    p.relative_basis='ttm';p.peers[0].reporting_basis='ttm';p.peers[0].financial_period='2025-Q4'
    r=evaluate(p)['models']['Trading comparables']
    assert next(v for v in r['ranges'] if v['multiple']=='ev_revenue')['enterprise_value']==4000000
    assert r['target']['values']['revenue']==400000
    book=load_workbook(BytesIO(export_workbench(p)))
    assert book['Comparables target']['B2'].value.count('SUMIFS')==4
    p.figures.pop()
    assert 'Trading comparables' in evaluate(p)['errors']

def test_forward_and_unverified_peer_bases_block_until_compatible():
    p=project();p.peers[0].reporting_basis='forward'
    assert 'incompatible' in evaluate(p)['errors']['Trading comparables']
    p.relative_basis='forward';p.forward_target=ForwardTarget(reviewed=True,currency='USD',period='2026',metrics={'revenue':200000})
    p.peers[0].financial_period='2026'
    assert evaluate(p)['models']['Trading comparables']['ranges'][0]['enterprise_value']==360000
    p.peers[0].basis_reviewed=False
    assert 'review' in evaluate(p)['errors']['Trading comparables']

def test_duplicate_selected_peers_and_distinct_transactions():
    p=project();copy=p.peers[0].model_copy();copy.company=' P.E.E.R '
    p.peers.append(copy)
    assert 'duplicate' in evaluate(p)['errors']['Trading comparables']
    p.peers[-1].selected=False
    stats=evaluate(p)['models']['Trading comparables']['statistics']['ev_revenue']
    assert stats['count']==1 and stats['median']==stats['average']==1.8
    p.peers[-1].selected=True;p.peers[-1].financial_period='2024'
    assert evaluate(p)['models']['Trading comparables']['statistics']['ev_revenue']['count']==2
    p.transactions.append(p.transactions[0].model_copy(update={'acquisition_date':'2026-02-01','measurement_date':'2026-02-01'}))
    assert evaluate(p)['models']['Precedent transactions']['statistics']['ev_revenue']['count']==2
    p.transactions.append(p.transactions[0].model_copy())
    assert 'duplicate' in evaluate(p)['errors']['Precedent transactions']

def test_short_fiscal_year_and_quarter_forecast_are_blocked():
    p=project();p.figures[0].period_start='2025-06-01'
    with pytest.raises(ValueError,match='fiscal year'): evaluate(p)
    p=project();p.figures=quarters()
    assert 'annual' in evaluate(p)['errors']['forecast'].lower()

def test_source_block_and_export_removed_audit():
    p=project();p.blocked_assumptions=['dcf.wacc']
    with pytest.raises(ValueError,match='Source document removed'): evaluate(p)
    with pytest.raises(ValueError,match='Source document removed'): export_workbench(p)

def test_fcf_ttm_uses_compatible_flows_and_calendar_changes_disable_growth():
    from equitylens.financials import analysis_table
    rows=quarters('operating_cash_flow')+quarters('capital_expenditure')
    for f in rows:
        if f.metric=='capital_expenditure':f.value=-20
    assert next(r for r in ttm_figures(rows,'USD')['figures'] if r['metric']=='free_cash_flow')['value']==320000
    rows=[FinancialFigure(company_id=1,metric='revenue',period=str(year),value=value,fiscal_calendar=calendar) for year,value,calendar in [(2024,100,'old'),(2025,200,'new')]]
    assert analysis_table(rows).loc['2025','revenue_growth'] is None

def test_import_metadata_is_conservative_and_legacy_comparable_api_rejects_duplicates():
    text='company,currency,unit_scale,ticker,financial_period,reporting_basis,measurement_date,period_end,basis_reviewed\nAcme,USD,1,,2025,annual,2026-01-01,2025-12-31,true'
    response=client.post('/v1/workbench/import-peers',files={'file':('peers.csv',text,'text/csv')})
    assert response.status_code==200
    row=response.json()['rows'][0]
    assert row['ticker']=='' and not row['selected'] and not row['basis_reviewed']
    row.update(selected=True,basis_reviewed=True,revenue=100,enterprise_value=1000)
    assert client.post('/v1/valuation/comparables',json={'rows':[row,row]}).status_code==422
    assert client.post('/v1/valuation/comparables',json={'rows':[row]}).status_code==200

def test_export_preserves_removed_source_and_analyst_resolution_history():
    from equitylens.assumption_export import AcceptedSource, source_rows
    from test_assumptions import extract
    p=project();source=extract('Revenue growth guidance 5% FY 2027')[0]
    original=AcceptedSource(source=source,target='forecast.growth',value=p.forecast.growth,period='2027',rationale='Reviewed source',classification='management guidance',accepted_at='2026-01-01',source_status='removed')
    with pytest.raises(ValueError,match='Source document removed'): export_workbench(p,[original])
    reviewed=original.model_copy(update={'source_status':'analyst_reviewed','classification':'analyst assumption','accepted_at':'2026-01-02','reviewed_at':'2026-01-02'})
    rows=source_rows(p,[original,reviewed])
    assert 'Superseded' in rows[0][4] and 'removed' in rows[0][4]
    assert 'User-reviewed analyst' in rows[1][4] and rows[1][6]=='analyst assumption'
    assert rows[0][22]=='2026-01-01' and rows[1][22]=='2026-01-02'
    book=load_workbook(BytesIO(export_workbench(p,[original,reviewed])))
    assert 'removed' in book['Assumptions & Sources']['E3'].value
    assert book['Assumptions & Sources']['U3'].value=='Source removed; text unavailable'
