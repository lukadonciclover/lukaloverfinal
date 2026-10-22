import pytest
from equitylens.historical_references import historical_references
from equitylens.models import FinancialFigure
from equitylens.reference_derivations import NWC_ASSETS, NWC_LIABILITIES

def f(metric,value,period='2025',notes=''):
    return FinancialFigure(company_id=1,metric=metric,value=value,period=period,notes=notes)

def average(rows,key):
    return next(a for a in historical_references(rows,'USD')['averages'] if a['metric']==key)

def test_reconciled_tax_and_reported_priority():
    rows=[f('net_income',75),f('tax_expense',25)]+[f(k,0) for k in ['discontinued_operations_net_income','noncontrolling_net_income','other_tax_reconciliation']]
    assert average(rows,'tax_rate')['value']==.25
    assert len(average(rows,'tax_rate')['source']['inputs'])==5
    assert average(rows+[f('pretax_income',125)],'tax_rate')['value']==.2
    rows[-1].value=5
    assert average(rows,'tax_rate')['value'] is None
    assert average(rows[:-1],'tax_rate')['value'] is None
    assert average([f('pretax_income',-10),f('tax_expense',2)],'tax_rate')['value'] is None
    assert average([f('pretax_income',10),f('tax_expense',20)],'tax_rate')['value'] is None

def test_ebit_bridge_requires_all_nonoperating_adjustments():
    rows=[f('ebit',100),f('interest_expense',10),f('nonoperating_expenses',20),f('nonoperating_income',10),f('tax_expense',20)]
    assert average(rows,'tax_rate')['value']==.25
    assert average(rows[:3]+rows[4:],'tax_rate')['value'] is None

def test_nwc_complete_components_and_source_provenance():
    rows=[f(k,10) for k in NWC_ASSETS]+[f(k,5) for k in NWC_LIABILITIES]+[f('revenue',100),f('cash',999),f('current_assets',1000),f('current_liabilities',5)]
    assert average(rows,'nwc_ratio')['value']==.25
    reference=next(a for a in historical_references(rows,'USD')['latest'] if a['metric']=='opening_nwc')
    assert reference['value']==25 and len(reference['inputs'])==7
    missing=average(rows[1:],'nwc_ratio')
    assert missing['value'] is None and 'accounts receivable' in missing['unavailable']
    assert average(rows+[f('net_working_capital',0)],'nwc_ratio')['value']==0

def test_borrowing_reference_requires_comparable_balances_and_definition():
    rows=[f('total_debt',100,'2024'),f('total_debt',120),f('interest_expense',5.5,notes='Borrowing interest')]
    assert average(rows,'interest_rate')['value']==.05
    rows[-1].notes='Net financing costs'
    assert average(rows,'interest_rate')['value'] is None
    rows[-1].notes='Borrowing interest'
    rows[1].value=200
    assert average(rows,'interest_rate')['value'] is None
    rows[1].value=0
    assert average(rows,'interest_rate')['value'] is None
    assert average(rows[:-1],'interest_rate')['value'] is None

def test_integrated_forecast_uses_same_complete_nwc_without_mutating_ledger():
    from test_workbench import project
    from equitylens.workbench import evaluate
    p=project()
    expected=evaluate(p)['models']['DCF']['equity_value']
    nwc=next(row for row in p.figures if row.metric=='net_working_capital')
    p.figures=[row for row in p.figures if row.metric!='net_working_capital']
    p.figures += [f(k,nwc.value if i==0 else 0,nwc.period) for i,k in enumerate((*NWC_ASSETS,*NWC_LIABILITIES))]
    # Match the source statement scale.
    for row in p.figures:
        if row.metric in (*NWC_ASSETS,*NWC_LIABILITIES): row.unit_scale=nwc.unit_scale
    assert evaluate(p)['models']['DCF']['equity_value']==pytest.approx(expected)
    assert not any(row.metric=='net_working_capital' for row in p.figures)
    from io import BytesIO
    from openpyxl import load_workbook
    from equitylens.workbook import export_workbench
    p.forecast.growth=.0422375293017343
    book=load_workbook(BytesIO(export_workbench(p)))
    opening=book['Opening balances']
    cells={row[0].value:row[1].value for row in opening.iter_rows(min_row=2)}
    assert cells['net_working_capital'].startswith('=B')
    assert cells['net_working_capital'].count('-')==3
    assert book['Forecast assumptions']['B3'].value==.0422375293017343
