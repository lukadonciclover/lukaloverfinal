"""Synthetic evidence fixtures; never loaded by the application."""
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app
from equitylens.assumptions import extract_assumptions, derive_assumptions, derive_revenue_guidance
from equitylens.documents import suggest_financial_figures
from equitylens.models import FinancialFigure
from equitylens.assumption_export import AcceptedSource, source_rows
from equitylens.workbench import Workbench
from equitylens.workbook import export_workbench

client = TestClient(app)


def extract(text):
    return extract_assumptions([{"page": 2, "text": text}], "fixture", "Synthetic guidance.pdf")


@pytest.mark.parametrize("text,metric,value,low,high", [
    ("Management expects revenue growth of 8% in FY 2027.", "revenue_growth", .08, None, None),
    ("Management expects revenue to increase between 5% and 7% next fiscal year.", "revenue_growth", None, .05, .07),
    ("FY 2027 EBITDA margin guidance 22%-24%.", "ebitda_margin", None, .22, .24),
    ("Management expects capital expenditure to represent approximately 4% of revenue.", "capex_ratio", .04, None, None),
    ("Capex budget USD 40 million in 2027.", "capex", 40, None, None),
    ("Expected revenue growth is strong next fiscal year.", "revenue_growth", None, None, None),
    ("WACC assumption 9%.", "wacc", .09, None, None),
    ("Dividend growth guidance 3% for 2027.", "dividend_growth", .03, None, None),
])
def test_explicit_and_qualitative(text, metric, value, low, high):
    a = extract(text)[0]
    assert (a["metric"], a["value"], a["minimum"], a["maximum"]) == (metric, value, low, high)
    assert a["page"] == 2 and a["excerpt"] and a["document_id"] == "fixture"
    if value is None and low is None:
        assert a["status"] == "requires review"


@pytest.mark.parametrize("word,scale", [("thousand", 1000), ("million", 1e6), ("billion", 1e9)])
def test_amount_scales(word, scale):
    a = extract(f"Consolidated revenue forecast USD 5 {word} in FY 2027.")[0]
    assert a["unit"] == "currency" and a["unit_scale"] == scale and a["currency"] == "USD"
    assert a["scope"] == "consolidated" and a["value"] == 5


@pytest.mark.parametrize("period,basis,expected", [("FY 2027","annual","FY 2027"),("calendar 2027","annual","CY 2027"),("Q4 2027","quarter","Q4 2027"),("2027-2029 CAGR","multi-year","2027 / 2029"),("long-term","long-term","Long term"),("next fiscal year","relative annual","next fiscal year")])
def test_periods(period,basis,expected):
    a=extract(f"Revenue growth guidance 5% {period}")[0]
    assert a["basis"]==basis and a["period"]==expected


def test_historical_guidance_and_scenarios_separate():
    text = "Revenue growth 4% FY 2025\nRevenue growth guidance 6% FY 2027\nIllustrative revenue growth 10% FY 2027"
    assert [a["classification"] for a in extract(text)] == ["historical reported","management guidance","illustrative scenario"]
    assert suggest_financial_figures([{"page":1,"text":text}]) == []
    assert suggest_financial_figures([{"page":1,"text":"Guidance FY 2027\nRevenue 500\nHistorical\nRevenue 2025 100"}])[0]["value"]==100


def test_structured_upload_and_no_ledger_contamination():
    data=b"metric,period,value,unit,currency,unit_scale,classification,scope\nrevenue_growth,FY 2027,6,percent,USD,1,management guidance,consolidated\nrevenue,2025,100,currency,USD,1000000,historical reported,consolidated\n"
    r=client.post("/v1/documents/process",files={"file":("Synthetic.csv",data,"text/csv")})
    assert r.status_code==200
    assert len(r.json()["suggestions"])==1
    assert r.json()["assumptions"][0]["value"]==.06


def figure(metric,value,period="2025",scale=1,currency="USD"):
    return FinancialFigure(company_id=1,metric=metric,value=value,period=period,unit_scale=scale,currency=currency)


def test_derived_only_valid_approved_inputs():
    fs=[figure("revenue",100,"2024",1e6),figure("revenue",110000,"2025",1000),figure("ebitda",22e6),figure("ebit",11e6),figure("capital_expenditure",-4.4e6),figure("depreciation_amortization",0),figure("net_working_capital",11e6),figure("tax_expense",2e6),figure("pretax_income",10e6)]
    rows={r["metric"]:r for r in derive_assumptions(fs,"USD")}
    assert rows["revenue_growth"]["value"]==pytest.approx(.1)
    assert rows["ebitda_margin"]["value"]==.2
    assert rows["capex_ratio"]["value"]==.04
    assert rows["depreciation_ratio"]["value"]==0
    assert rows["tax_rate"]["value"]==.2
    assert "not a cash" in rows["tax_rate"]["rationale"]
    assert all(r["classification"]=="historical derived" for r in rows.values())


@pytest.mark.parametrize("denominator", [None,0,-1])
def test_invalid_denominators(denominator):
    assert derive_assumptions([figure("revenue",denominator),figure("ebitda",5)],"USD")==[]


def test_duplicate_currency_quarter_not_derived():
    assert derive_assumptions([figure("revenue",100),figure("revenue",200),figure("ebitda",10)],"USD")==[]
    assert derive_assumptions([figure("revenue",100,currency="EUR"),figure("ebitda",10)],"USD")==[]
    assert derive_assumptions([figure("revenue",100,"2025-Q1"),figure("ebitda",10,"2025-Q1")],"USD")==[]


def test_absolute_revenue_derivation_and_missing_base():
    a=extract("Consolidated revenue forecast USD 110 million in FY 2026.")
    fs=[figure("revenue",100,scale=1e6)]
    result=derive_revenue_guidance(a,fs,"USD")
    assert result[0]["value"]==pytest.approx(.1)
    assert result[0]["original"]["value"]==110
    assert derive_revenue_guidance(a,[],"USD")==[]
    assert derive_revenue_guidance(a,[figure("revenue",0)],"USD")==[]
    assert derive_revenue_guidance(a,fs,"EUR")==[]


def test_no_text_no_guidance():
    assert extract("")==[]
    assert extract("Auditor signature and company address")==[]
    assert extract("Revenue 2025 100")==[]


def test_wide_forecast_columns_and_xlsx_upload():
    from openpyxl import Workbook
    wb=Workbook()
    sheet=wb.active
    sheet.title="Guidance"
    sheet.append(["metric","FY 2027 forecast","FY 2028 forecast"])
    sheet.append(["revenue_growth","5%-7%","8%"])
    buffer=BytesIO();wb.save(buffer)
    response=client.post("/v1/documents/process",files={"file":("Synthetic.xlsx",buffer.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code==200
    rows=response.json()["assumptions"]
    assert len(rows)==2
    assert rows[0]["minimum"]==.05 and rows[0]["maximum"]==.07 and rows[0]["value"] is None
    assert rows[1]["value"]==.08
    assert response.json()["suggestions"]==[]


def test_pdf_and_scanned_pdf_processing():
    import pymupdf
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((50,50),"Consolidated revenue growth guidance 6% for FY 2027")
    response=client.post("/v1/documents/process",files={"file":("Synthetic.pdf",pdf.tobytes(),"application/pdf")})
    assert response.status_code==200 and response.json()["assumptions"][0]["value"]==.06
    pdf.close()
    blank=pymupdf.open();blank.new_page()
    response=client.post("/v1/documents/process",files={"file":("Synthetic-scan.pdf",blank.tobytes(),"application/pdf")})
    blank.close()
    assert response.status_code==200 and response.json()["assumptions"]==[] and "OCR" in response.json()["document"]["warning"]


def test_accepted_export_and_effective_value():
    from test_workbench import model
    a=extract("Consolidated revenue growth guidance 5% FY 2027")[0]
    accepted=AcceptedSource(source=a,target="dcf.revenue_growth.0",value=.05,period="2027",rationale="Reviewed source",classification="management guidance",accepted_at="2026-10-22T00:00:00Z")
    p=Workbench(currency="USD",dcf=model(),cash=0,debt=0)
    p.dcf.revenue_growth[0]=.05
    rows=source_rows(p,[accepted])
    assert rows[0][3]==.05 and rows[0][4]=="Current model input"
    wb=load_workbook(BytesIO(export_workbench(p,[accepted])))
    assert wb["Assumptions & Sources"]["B2"].value=="dcf.revenue_growth.0"
    assert wb["Assumptions & Sources"]["C2"].value==.05
    assert wb.calculation.fullCalcOnLoad
    p.dcf.revenue_growth[0]=.07
    assert "Superseded" in source_rows(p,[accepted])[0][4]
    p.dcf=None
    assert "disabled" in source_rows(p,[accepted])[0][4]


def test_dividend_growth_is_per_share_not_earnings():
    fs=[figure("dividend_per_share",2,"2024"),figure("dividend_per_share",2.2,"2025"),figure("net_income",500)]
    rows=derive_assumptions(fs,"USD")
    assert len(rows)==1 and rows[0]["metric"]=="dividend_growth" and rows[0]["value"]==pytest.approx(.1)


def test_calculate_and_export_api_with_accepted_input_and_redacted_source():
    from test_workbench import model
    d=model();d.revenue_growth[0]=.06
    payload=Workbench(currency="USD",dcf=d,cash=0,debt=0).model_dump(mode="json")
    calculation=client.post("/v1/workbench/calculate",json=payload)
    assert calculation.status_code==200 and "DCF" in calculation.json()["models"]
    a=extract("Revenue growth guidance 6% FY 2027")[0];a["excerpt"]=""
    payload["assumption_sources"]=[dict(source=a,target="dcf.revenue_growth.0",value=.06,period="2027",rationale="Confirmed annual input",classification="management guidance",accepted_at="2026-10-22T00:00:00Z")]
    response=client.post("/v1/workbench/export",json=payload)
    assert response.status_code==200
    wb=load_workbook(BytesIO(response.content))
    audit=wb["Assumptions & Sources"]
    assert audit["D2"].value==.06 and audit["U2"].value=="Source text unavailable"
    assert audit["X2"].value=="Accepted"


def test_effective_dcf_export_does_not_misattribute_overridden_input():
    from test_workbench import project
    p=project()
    a=extract("Revenue growth guidance 80% FY 2027")[0]
    accepted=AcceptedSource(source=a,target="dcf.revenue_growth.0",value=.8,period="2027",rationale="Old standalone input",classification="analyst assumption",accepted_at="2026-10-22T00:00:00Z")
    p.dcf.revenue_growth[0]=.8
    row=source_rows(p,[accepted])[0]
    assert row[3]==p.forecast.growth and "Superseded" in row[4]
