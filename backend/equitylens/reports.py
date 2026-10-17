from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

import pandas as pd
from docx import Document
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def thesis_lines(thesis: dict[str, list[dict[str, str]]]) -> list[tuple[str, list[str]]]:
    return [
        (heading, [f"[{item['classification']}] {item['text']}" for item in items] or ["Insufficient verified information."])
        for heading, items in thesis.items()
    ]


def generate_docx(company: Any, thesis: dict[str, list[dict[str, str]]], figures: pd.DataFrame, valuation: dict[str, Any] | None) -> bytes:
    document = Document()
    document.add_heading(f"{company.name} Investment Research", 0)
    document.add_paragraph(f"As of {date.today().isoformat()} | Currency: {company.currency} | EquityLens AI")
    if getattr(company, "is_synthetic", False):
        document.add_paragraph("SYNTHETIC DEMONSTRATION DATA - NOT INVESTMENT RESEARCH")
    document.add_page_break()
    for heading, lines in thesis_lines(thesis):
        document.add_heading(heading, level=1)
        for line in lines:
            document.add_paragraph(line)
    document.add_heading("Historical financial data", level=1)
    _add_docx_table(document, figures)
    if valuation:
        document.add_heading("DCF model results", level=1)
        for key in ("enterprise_value", "equity_value", "implied_share_value", "terminal_value"):
            if valuation.get(key) is not None:
                document.add_paragraph(f"{key.replace('_', ' ').title()}: {valuation[key]:,.2f}")
        document.add_paragraph(valuation.get("timing_convention", ""))
    document.add_heading("Sources and assumptions", level=1)
    document.add_paragraph("Figures retain their origin and source metadata in the application. User-entered assumptions are not reported facts.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_docx_table(document: Document, frame: pd.DataFrame) -> None:
    if frame.empty:
        document.add_paragraph("No verified historical figures available.")
        return
    display = frame[[column for column in ["metric", "period", "value", "currency", "origin"] if column in frame.columns]]
    table = document.add_table(rows=1, cols=len(display.columns))
    for index, column in enumerate(display.columns):
        table.rows[0].cells[index].text = str(column).replace("_", " ").title()
    for _, row in display.iterrows():
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = "" if pd.isna(value) else str(value)


def generate_pdf(company: Any, thesis: dict[str, list[dict[str, str]]], valuation: dict[str, Any] | None) -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    story: list[Any] = [
        Paragraph(f"{company.name} Investment Research", styles["Title"]),
        Spacer(1, 0.2 * inch),
        Paragraph(f"As of {date.today().isoformat()} | Currency: {company.currency}", styles["Normal"]),
    ]
    if getattr(company, "is_synthetic", False):
        story.append(Paragraph("SYNTHETIC DEMONSTRATION DATA - NOT INVESTMENT RESEARCH", styles["Heading2"]))
    story.append(PageBreak())
    for heading, lines in thesis_lines(thesis):
        story.append(Paragraph(heading, styles["Heading1"]))
        for line in lines:
            story.append(Paragraph(line, styles["BodyText"]))
            story.append(Spacer(1, 0.08 * inch))
    if valuation:
        story.append(Paragraph("DCF model results", styles["Heading1"]))
        rows = [["Measure", "Value"]]
        for key in ("enterprise_value", "equity_value", "implied_share_value", "terminal_value"):
            if valuation.get(key) is not None:
                rows.append([key.replace("_", " ").title(), f"{valuation[key]:,.2f}"])
        table = Table(rows, colWidths=[3 * inch, 2 * inch])
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183B4E")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]))
        story.append(table)
    SimpleDocTemplate(output, pagesize=letter, title=f"{company.name} Investment Research").build(story)
    return output.getvalue()


def generate_excel(company: Any, figures: pd.DataFrame, assumptions: pd.DataFrame, forecast: pd.DataFrame, valuation: dict[str, Any] | None) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame({"Field": ["Company", "As of date", "Currency", "Synthetic"], "Value": [company.name, date.today().isoformat(), company.currency, company.is_synthetic]}).to_excel(writer, sheet_name="Overview", index=False)
        figures.to_excel(writer, sheet_name="Historical Financials", index=False)
        assumptions.to_excel(writer, sheet_name="DCF Assumptions", index=False)
        forecast.to_excel(writer, sheet_name="DCF Forecast")
        pd.DataFrame([valuation or {}]).to_excel(writer, sheet_name="Valuation Summary", index=False)
        for sheet in writer.book.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="183B4E")
            sheet.freeze_panes = "A2"
    return output.getvalue()
