"""Excel export using the application's existing openpyxl dependency."""

from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from equitylens.workbench import Workbench, evaluate, effective_dcf


def export_workbench(p: Workbench, assumption_sources=None) -> bytes:
    result = evaluate(p)
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    summary = wb.create_sheet("Valuation summary")
    summary.append(["Method", "Equity value", "Per share", "Status / basis"])
    bridge = wb.create_sheet("Equity bridge")
    bridge.append(["Input", "Value"])
    for k in ("cash", "debt", "minority", "non_operating_assets", "shares", "price"):
        bridge.append([k, getattr(p, k)])

    def safe(value):
        return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value

    sources = wb.create_sheet("Sources")
    sources.append(["Currency", p.currency, "All money in actual units; shares in actual shares"])
    sources.append(
        [
            "Methodology",
            "Year-end discounting. No tax benefit on losses unless selected. Peer ranges are observed min/max, not confidence intervals.",
        ]
    )
    sources.append(
        [
            "Forecast policy",
            "NWC is net operating assets. Other assets/liabilities held constant. Interest on opening debt. Revolver drawn at year-end to maintain minimum cash.",
        ]
    )
    sources.append(
        [
            "Model differences",
            "DCF values operating cash flows; DDM values shareholder dividends; NAV values adjusted assets. Multiples reflect selected market/transaction observations.",
        ]
    )
    sources.append(["Input snapshot", json.dumps(p.model_dump(mode="json"))[:32000]])
    history = wb.create_sheet("Historical statements")
    history.append(
        [
            "Fiscal period",
            "Metric",
            "Reported value",
            "Currency",
            "Unit scale",
            "Actual units",
            "Source",
            "Page",
            "Notes",
            "Source excerpt",
        ]
    )
    for f in p.figures:
        r = history.max_row + 1
        history.append(
            [
                safe(f.period),
                safe(f.metric),
                f.value,
                f.currency,
                f.unit_scale,
                f'=IF(C{r}="","",C{r}*E{r})',
                safe(f.source.document_name) if f.source else "Manual input",
                f.source.page if f.source else None,
                safe(f.notes),
                safe(f.source.quote) if f.source else "",
            ]
        )
    assumptions = wb.create_sheet("Forecast assumptions")
    assumptions.append(["Assumption", "Value"])
    if p.forecast:
        for k, v in p.forecast.model_dump().items():
            assumptions.append(["Exported horizon (fixed; re-export to change)" if k == "horizon" else k, v])
    else:
        assumptions.append(["Status", "Not completed"])
    projected = wb.create_sheet("Projected statements")
    if result["forecast"]:
        opening = wb.create_sheet("Opening balances")
        opening.append(["Metric", "Actual units", "Fiscal year"])
        period = max(result["history"])
        opening_rows = {}
        for key in result["history"][period]:
            r = opening.max_row + 1
            opening_rows[key] = r
            opening.append(
                [
                    key,
                    f"=SUMIFS('Historical statements'!F2:F{history.max_row},'Historical statements'!A2:A{history.max_row},C{r},'Historical statements'!B2:B{history.max_row},A{r})",
                    period,
                ]
            )

        def op(key):
            return f"'Opening balances'!$B${opening_rows[key]}"

        keys = list(result["forecast"][0])
        projected.append(keys)
        cols = {k: get_column_letter(i + 1) for i, k in enumerate(keys)}
        for i in range(p.forecast.horizon):
            r = i + 2

            def cell(k):
                return f"{cols[k]}{r}"

            def prev(k):
                return op(k) if i == 0 else f"{cols[k]}{r - 1}"

            debt_prev = f"({op('short_term_debt')}+{op('long_term_debt')})" if i == 0 else prev("debt")
            a = lambda row: f"'Forecast assumptions'!$B${row}"
            formulas = {
                "year": str(int(period) + i + 1),
                "revenue": f"{prev('revenue')}*(1+{a(3)})",
                "gross_profit": f"{cell('revenue')}*{a(4)}",
                "cost_of_goods_sold": f"{cell('revenue')}-{cell('gross_profit')}",
                "operating_expenses": f"{cell('revenue')}*IF(ISNUMBER({a(6)}),{a(6)},{a(4)}-{a(5)})",
                "ebitda": f"{cell('gross_profit')}-{cell('operating_expenses')}",
                "depreciation": f"{cell('revenue')}*{a(7)}",
                "ebit": f"{cell('ebitda')}-{cell('depreciation')}",
                "interest_expense": f"{debt_prev}*{a(11)}",
                "tax_expense": f"MAX(0,{cell('ebit')}-{cell('interest_expense')})*{a(8)}",
                "net_income": f"{cell('ebit')}-{cell('interest_expense')}-{cell('tax_expense')}",
                "net_working_capital": f"{cell('revenue')}*{a(10)}",
                "change_nwc": f"{cell('net_working_capital')}-{prev('net_working_capital')}",
                "operating_cash_flow": f"{cell('net_income')}+{cell('depreciation')}-{cell('change_nwc')}",
                "capital_expenditure": f"{cell('revenue')}*{a(9)}",
                "investing_cash_flow": f"-{cell('capital_expenditure')}",
                "revolver_draw": f'IF({a(15)}="revolver",MAX(0,{a(14)}-({prev("cash")}+{cell("operating_cash_flow")}-{cell("capital_expenditure")}+{a(12)}-{a(13)})),0)',
                "financing_cash_flow": f"{a(12)}+{cell('revolver_draw')}-{a(13)}",
                "cash": f"{prev('cash')}+{cell('operating_cash_flow')}+{cell('investing_cash_flow')}+{cell('financing_cash_flow')}",
                "debt": f"{debt_prev}+{a(12)}+{cell('revolver_draw')}",
                "property_plant_equipment": f"{prev('property_plant_equipment')}+{cell('capital_expenditure')}-{cell('depreciation')}",
                "other_assets": f"{op('total_assets')}-{op('cash')}-{op('net_working_capital')}-{op('property_plant_equipment')}",
                "other_liabilities": f"{op('total_liabilities')}-{op('short_term_debt')}-{op('long_term_debt')}",
                "total_assets": f"{cell('cash')}+{cell('net_working_capital')}+{cell('property_plant_equipment')}+{cell('other_assets')}",
                "total_liabilities": f"{cell('debt')}+{cell('other_liabilities')}",
                "total_equity": f"{prev('total_equity')}+{cell('net_income')}-{a(13)}",
                "balance_check": f"{cell('total_assets')}-{cell('total_liabilities')}-{cell('total_equity')}",
                "free_cash_flow": f"{cell('operating_cash_flow')}-{cell('capital_expenditure')}",
                "fcff": f"{cell('ebit')}-MAX(0,{cell('ebit')})*{a(8)}+{cell('depreciation')}-{cell('capital_expenditure')}-{cell('change_nwc')}",
            }
            projected.append(["=" + formulas[k] for k in keys])
    else:
        projected.append(["Not completed", result["errors"].get("forecast", "No forecast selected")])
    wacc = wb.create_sheet("WACC")
    wacc.append(["Assumption", "Value"])
    if p.wacc:
        for k, v in p.wacc.model_dump().items():
            wacc.append([k, v])
        wacc.append(["Cost of equity", "=B2+B3*B4"])
        wacc.append(["WACC", "=IF(B7+B8>0,B9*B7/(B7+B8)+B5*(1-B6)*B8/(B7+B8),NA())"])
    else:
        wacc.append(["Manual WACC", p.dcf.wacc if p.dcf else None])
    dcf_inputs = wb.create_sheet("DCF inputs")
    dcf_inputs.append(["Input", "Value"])
    base = effective_dcf(p, result["history"], result["forecast"]) if "DCF" in result["models"] else None
    input_rows = {}
    if base:
        for k, v in base.model_dump().items():
            if not isinstance(v, list):
                input_rows[k] = dcf_inputs.max_row + 1
                dcf_inputs.append([k, v])
        dcf_inputs.cell(input_rows["wacc"], 2, "='WACC'!B10" if p.wacc else "='WACC'!B2")
        for key, row in [
            ("cash", 2),
            ("debt", 3),
            ("minority_interest", 4),
            ("non_operating_assets", 5),
            ("diluted_shares", 6),
            ("reference_price", 7),
        ]:
            dcf_inputs.cell(input_rows[key], 2, f"='Equity bridge'!B{row}")
        if p.forecast:
            dcf_inputs.cell(input_rows["base_revenue"], 2, "=" + op("revenue"))
            dcf_inputs.cell(input_rows["opening_nwc"], 2, "=" + op("net_working_capital"))

    def ref(key):
        return f"'DCF inputs'!$B${input_rows[key]}"

    def build_dcf(name, growth_delta=0, margin_delta=0, wacc_delta=0, terminal_delta=0, exit_delta=0, method=None):
        sheet = wb.create_sheet(name)
        sheet.append(
            [
                "Year",
                "Growth",
                "EBIT margin",
                "Tax rate",
                "D&A ratio",
                "Capex ratio",
                "NWC ratio",
                "Revenue",
                "EBIT",
                "Cash taxes",
                "NOPAT",
                "D&A",
                "Capex",
                "NWC",
                "Change NWC",
                "FCFF",
                "PV FCFF",
            ]
        )
        sheet["S1"] = "WACC"
        sheet["T1"] = f"={ref('wacc')}+({wacc_delta})"
        sheet["S2"] = "Terminal growth"
        sheet["T2"] = f"={ref('terminal_growth')}+({terminal_delta})"
        sheet["S3"] = "Exit multiple"
        sheet["T3"] = f"={ref('exit_multiple')}+({exit_delta})"
        keys = [
            "revenue_growth",
            "ebit_margin",
            "tax_rate",
            "depreciation_pct_revenue",
            "capex_pct_revenue",
            "nwc_pct_revenue",
        ]
        for i in range(len(base.revenue_growth)):
            r = i + 2
            sheet.cell(r, 1, i + 1)
            for c, k in enumerate(keys, 2):
                if name == "DCF":
                    linked = {
                        2: "='Forecast assumptions'!B3",
                        3: "=IF(ISNUMBER('Forecast assumptions'!B6),'Forecast assumptions'!B4-'Forecast assumptions'!B6,'Forecast assumptions'!B5)-'Forecast assumptions'!B7",
                        4: "='Forecast assumptions'!B8",
                        5: "='Forecast assumptions'!B7",
                        6: "='Forecast assumptions'!B9",
                        7: "='Forecast assumptions'!B10",
                    }
                    sheet.cell(r, c, linked[c] if p.forecast else getattr(base, k)[i])
                else:
                    delta = growth_delta if c == 2 else margin_delta if c == 3 else 0
                    sheet.cell(r, c, f"='DCF'!{get_column_letter(c)}{r}+({delta})")
            previous = ref("base_revenue") if i == 0 else f"H{r - 1}"
            nwc_previous = (
                (ref("opening_nwc") if base.opening_nwc is not None else f"{ref('base_revenue')}*G2")
                if i == 0
                else f"N{r - 1}"
            )
            formulas = [
                f"{previous}*(1+B{r})",
                f"H{r}*C{r}",
                f"IF({ref('tax_loss_benefit')},I{r},MAX(0,I{r}))*D{r}",
                f"I{r}-J{r}",
                f"H{r}*E{r}",
                f"H{r}*F{r}",
                f"H{r}*G{r}",
                f"N{r}-({nwc_previous})",
                f"K{r}+L{r}-M{r}-O{r}",
                f"P{r}/(1+$T$1)^A{r}",
            ]
            for c, formula in enumerate(formulas, 8):
                sheet.cell(r, c, "=" + formula)
        n = len(base.revenue_growth) + 1
        growth_terminal = f"IF(AND(T1>0,T1>T2),P{n}*(1+T2)/(T1-T2),NA())"
        exit_terminal = f"IF(AND(T1>0,T3>0,I{n}+L{n}>0),(I{n}+L{n})*T3,NA())"
        terminal = (
            (growth_terminal if method == "growth" else exit_terminal)
            if method
            else f'IF({ref("terminal_method")}="growth",{growth_terminal},{exit_terminal})'
        )
        for r, label, formula in [
            (5, "Terminal value", terminal),
            (6, "PV terminal", f"T5/(1+T1)^A{n}"),
            (7, "Enterprise value", f"SUM(Q2:Q{n})+T6"),
            (
                8,
                "Equity value",
                f"T7+{ref('cash')}-{ref('debt')}-{ref('minority_interest')}+{ref('non_operating_assets')}",
            ),
            (9, "Per share", f'IF({ref("diluted_shares")}>0,T8/{ref("diluted_shares")},"")'),
        ]:
            sheet.cell(r, 19, label)
            sheet.cell(r, 20, "=" + formula)
        return sheet

    if base:
        build_dcf("DCF")
        summary.append(["DCF", "='DCF'!T8", "='DCF'!T9", "Operating cash flows / enterprise-to-equity bridge"])
    else:
        wb.create_sheet("DCF").append(["Not completed", result["errors"].get("dcf", "No DCF selected")])
    scenarios = wb.create_sheet("Scenarios")
    scenarios.append(
        [
            "Scenario",
            "Growth delta",
            "Margin delta",
            "WACC delta",
            "Terminal delta",
            "Exit delta",
            "Equity value",
            "Per share",
        ]
    )
    if base:
        for i, (name, s) in enumerate(p.scenarios.items(), 2):
            scenarios.append([name, *s.model_dump().values()])
            title = f"Case {i - 1}"
            build_dcf(title, *[f"'Scenarios'!{c}{i}" for c in "BCDEF"])
            scenarios.cell(i, 7, f"='{title}'!T8")
            scenarios.cell(i, 8, f"='{title}'!T9")
    sensitivity = wb.create_sheet("Sensitivity")
    sensitivity.append(["Type", "Column assumption", "Row assumption", "Equity value (live)"])
    if base:
        for kind, grid in result["sensitivities"].items():
            for iy, y in enumerate(grid["rows"]):
                for ix, x in enumerate(grid["columns"]):
                    r = sensitivity.max_row + 1
                    sensitivity.append([kind, x, y])
                    title = f"Sens {r - 1}"
                    if kind == "operating":
                        build_dcf(title, f"'Sensitivity'!B{r}", f"'Sensitivity'!C{r}")
                    else:
                        build_dcf(
                            title,
                            wacc_delta=f"'Sensitivity'!B{r}-{ref('wacc')}",
                            terminal_delta=f"'Sensitivity'!C{r}-{ref('terminal_growth')}" if kind == "growth" else 0,
                            exit_delta=f"'Sensitivity'!C{r}-{ref('exit_multiple')}" if kind == "exit_multiple" else 0,
                            method=kind,
                        )
                    wb[title].sheet_state = "hidden"
                    sensitivity.cell(r, 4, f"=IFERROR('{title}'!T8,\"Invalid assumptions\")")
    for title, rows, model in [
        ("Comparables", p.peers, "Trading comparables"),
        ("Transactions", p.transactions, "Precedent transactions"),
    ]:
        sheet = wb.create_sheet(title)
        sheet.append(
            [
                "Company",
                "Selected",
                "Currency",
                "Scale",
                "Market cap",
                "Debt",
                "Cash",
                "EV",
                "Revenue",
                "EBITDA",
                "EBIT",
                "Net income",
                "Book equity",
                "EV Revenue",
                "EV EBITDA",
                "EV EBIT",
                "PE",
                "PB",
                "PS",
                "Acquisition date",
            ]
        )
        for peer in rows:
            r = sheet.max_row + 1
            sheet.append(
                [
                    safe(peer.company),
                    peer.selected,
                    peer.currency,
                    peer.unit_scale,
                    peer.market_cap,
                    peer.debt,
                    peer.cash,
                    peer.enterprise_value,
                    peer.revenue,
                    peer.ebitda,
                    peer.ebit,
                    peer.net_income,
                    peer.book_equity,
                ]
            )
            if peer.enterprise_value is None:
                sheet.cell(r, 8, f'=IF(COUNT(E{r}:G{r})=3,E{r}+F{r}-G{r},"")')
            for col, num, den in [
                (14, "H", "I"),
                (15, "H", "J"),
                (16, "H", "K"),
                (17, "E", "L"),
                (18, "E", "M"),
                (19, "E", "I"),
            ]:
                sheet.cell(
                    r,
                    col,
                    f"=IF(AND(B{r},C{r}='Sources'!$B$1,{den}{r}>0,ISNUMBER({num}{r})),{num}{r}/{den}{r},\"\")"
                    if title == "Comparables" or col < 17
                    else None,
                )
            sheet.cell(r, 20, peer.acquisition_date)
        end = sheet.max_row
        stat_rows = {}
        if rows:
            for label, fn in [("Median", "MEDIAN"), ("Average", "AVERAGE"), ("Minimum", "MIN"), ("Maximum", "MAX")]:
                r = sheet.max_row + 1
                sheet.cell(r, 1, label)
                stat_rows[label] = r
                for c in range(14, 20):
                    letter = get_column_letter(c)
                    sheet.cell(r, c, f'=IF(COUNT({letter}2:{letter}{end})>0,{fn}({letter}2:{letter}{end}),"")')
        else:
            sheet.append(["Not completed"])
        relative = result["models"].get(model, {})
        if relative.get("ranges"):
            sheet.append(
                [
                    "Implied valuations",
                    "Target metric (actual units)",
                    "Low equity",
                    "Median equity",
                    "High equity",
                    "Per share",
                    "Median enterprise",
                ]
            )
            metric_map = {
                "ev_revenue": ("revenue", "N"),
                "ev_ebitda": ("ebitda", "O"),
                "ev_ebit": ("ebit", "P"),
                "pe": ("net_income", "Q"),
                "pb": ("total_equity", "R"),
                "ps": ("revenue", "S"),
            }
            period = max(result["history"])
            for val in relative["ranges"]:
                metric, col = metric_map[val["multiple"]]
                r = sheet.max_row + 1
                sheet.append(
                    [
                        val["multiple"],
                        f"=SUMIFS('Historical statements'!F2:F{history.max_row},'Historical statements'!A2:A{history.max_row},\"{period}\",'Historical statements'!B2:B{history.max_row},\"{metric}\")",
                    ]
                )
                adjustment = (
                    "+'Equity bridge'!B2-'Equity bridge'!B3-'Equity bridge'!B4+'Equity bridge'!B5"
                    if val["multiple"].startswith("ev_")
                    else ""
                )
                for c, stat in [(3, "Minimum"), (4, "Median"), (5, "Maximum")]:
                    sheet.cell(
                        r,
                        c,
                        f"=IF(AND(B{r}>0,ISNUMBER({col}{stat_rows[stat]})),B{r}*{col}{stat_rows[stat]}{adjustment},NA())",
                    )
                sheet.cell(r, 6, f"=IF('Equity bridge'!B6>0,D{r}/'Equity bridge'!B6,\"\")")
                if val["multiple"].startswith("ev_"):
                    sheet.cell(r, 7, f"=B{r}*{col}{stat_rows['Median']}")
                summary.append(
                    [
                        model + " / " + val["multiple"],
                        f"='{title}'!D{r}",
                        f"='{title}'!F{r}",
                        "Median selected multiple; see model sheet for min/max equity range",
                    ]
                )
    ddm = wb.create_sheet("DDM")
    ddm.append(["Input", "Value"])
    if p.ddm and "DDM" in result["models"]:
        for k, v in p.ddm.model_dump().items():
            ddm.append([k, v])
        ddm.append(["Year", "Dividend", "PV"])
        for i in range(1, 11):
            r = ddm.max_row + 1
            ddm.append([i, f"=IF(A{r}<=$B$6,$B$2*(1+$B$3)^A{r},0)", f"=B{r}/(1+$B$5)^A{r}"])
        last = ddm.max_row
        ddm["E1"] = "Per share"
        ddm["F1"] = (
            f"=IF(B6=0,IF(B5>B3,B2*(1+B3)/(B5-B3),NA()),IF(B5>B4,SUM(C8:C{last})+B2*(1+B3)^B6*(1+B4)/(B5-B4)/(1+B5)^B6,NA()))"
        )
        summary.append(
            [
                "DDM",
                "=IF('Equity bridge'!B6>0,'DDM'!F1*'Equity bridge'!B6,\"\")",
                "='DDM'!F1",
                "Dividend cash flows to equity",
            ]
        )
    else:
        ddm.append(["Not completed", result["errors"].get("DDM", "")])
    assets = wb.create_sheet("Asset valuation")
    assets.append(["Input", "Value"])
    if p.assets and "Asset based" in result["models"]:
        for k, v in p.assets.model_dump().items():
            assets.append([k, v])
        assets.append(["Diluted shares", "='Equity bridge'!B6"])
        assets.append(["Book equity", "=B2-B3"])
        assets.append(["Adjusted equity", "=IF(AND(B2+B4>=0,B3+B5>=0),B2+B4-B3-B5,NA())"])
        assets.append(["Per share", '=IF(B6>0,B8/B6,"")'])
        summary.append(
            [
                "Asset based",
                "='Asset valuation'!B8",
                "='Asset valuation'!B9",
                "Adjusted asset values less adjusted liabilities",
            ]
        )
    else:
        assets.append(["Not completed", result["errors"].get("Asset based", "")])
    for key, error in result["errors"].items():
        summary.append([key, None, None, error])
    from equitylens.assumption_export import source_rows
    audit = wb.create_sheet("Assumptions & Sources")
    audit.append(["Metric", "Target", "Accepted value", "Current model value", "Usage", "Application period", "Applied classification", "Original value", "Original minimum", "Original maximum", "Source unit", "Source currency", "Source scale", "Source period", "Period basis", "Scope", "Source classification", "Document ID", "Document", "Page", "Excerpt", "Rationale", "Accepted at", "Review status"])
    for row in source_rows(p, assumption_sources or [], result):
        audit.append([safe(value) for value in row])
    for sheet in wb:
        sheet.cell(
            sheet.max_row + 2, 1, f"{p.currency} | actual currency units; shares are actual counts; rates are decimals."
        )
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                formula = cell.data_type == "f"
                if isinstance(cell.value, str) and not formula and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
                cell.font = Font(
                    name="Calibri",
                    size=11,
                    color=("008000" if "!" in cell.value else "000000")
                    if formula
                    else "0000FF"
                    if isinstance(cell.value, (int, float))
                    else "243746",
                )
                if isinstance(cell.value, (int, float)) or formula:
                    cell.number_format = '#,##0.00;[Red](#,##0.00);"-"'
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="122B3A")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 24 if col[0].column > 1 else 32
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
